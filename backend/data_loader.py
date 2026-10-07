import os
import sys
import csv
from datetime import datetime, time
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from database import SessionLocal, engine, Base
from models import User, Station, Train, RouteStop, LiveTrainState, FieldObservation, ETAPrediction
import auth

DATA_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "data"))

# Reference Coordinates for stations along Northern Corridor
STATION_COORDS = {
    "NDLS": {"lat": 28.6427, "lng": 77.2195, "name": "New Delhi", "zone": "NR", "div": "DLI"},
    "SNP": {"lat": 28.9944, "lng": 77.0194, "name": "Sonipat Junction", "zone": "NR", "div": "DLI"},
    "PNP": {"lat": 29.3909, "lng": 76.9635, "name": "Panipat Junction", "zone": "NR", "div": "DLI"},
    "KUN": {"lat": 29.6857, "lng": 76.9905, "name": "Karnal", "zone": "NR", "div": "DLI"},
    "KKDE": {"lat": 29.9695, "lng": 76.8783, "name": "Kurukshetra Junction", "zone": "NR", "div": "DLI"},
    "UMB": {"lat": 30.3340, "lng": 76.8329, "name": "Ambala Cantt Junction", "zone": "NR", "div": "UMB"},
    "CDG": {"lat": 30.7046, "lng": 76.8206, "name": "Chandigarh Junction", "zone": "NR", "div": "UMB"},
    "KLK": {"lat": 30.8359, "lng": 76.9360, "name": "Kalka", "zone": "NR", "div": "UMB"},
    "LDH": {"lat": 30.9010, "lng": 75.8573, "name": "Ludhiana Junction", "zone": "NR", "div": "FZR"},
    "PGW": {"lat": 31.2223, "lng": 75.7699, "name": "Phagwara Junction", "zone": "NR", "div": "FZR"},
    "JUC": {"lat": 31.3260, "lng": 75.5762, "name": "Jalandhar City", "zone": "NR", "div": "FZR"},
    "BEAS": {"lat": 31.5160, "lng": 75.3050, "name": "Beas Junction", "zone": "NR", "div": "FZR"},
    "ASR": {"lat": 31.6340, "lng": 74.8723, "name": "Amritsar Junction", "zone": "NR", "div": "FZR"},
    "RE": {"lat": 28.1920, "lng": 76.6239, "name": "Rewari Junction", "zone": "NWR", "div": "JP"},
    "FKA": {"lat": 30.4042, "lng": 74.0252, "name": "Fazilka Junction", "zone": "NR", "div": "FZR"},
    "BTI": {"lat": 30.2110, "lng": 74.9455, "name": "Bathinda Junction", "zone": "NR", "div": "UMB"},
}

def clean_time(val):
    if val is None or val == "--" or val == "-":
        return None
    if isinstance(val, time):
        return val.strftime("%H:%M")
    if isinstance(val, datetime):
        return val.strftime("%H:%M")
    val_str = str(val).strip()
    if ":" in val_str:
        return val_str[:5]
    return None

def load_all_data():
    print("Initializing database tables...")
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()

    try:
        # 1. Seed Users if not present
        if not db.query(User).filter(User.email == "passenger@railnirdeshak.in").first():
            db.add_all([
                User(
                    email="passenger@railnirdeshak.in",
                    full_name="Rajesh Sharma",
                    hashed_password=auth.hash_password("demo123"),
                    role="passenger"
                ),
                User(
                    email="control@railnirdeshak.in",
                    full_name="Northern Zone Traffic Controller",
                    hashed_password=auth.hash_password("admin123"),
                    role="controller"
                )
            ])
            db.commit()
            print("✓ Default users seeded.")

        # 2. Extract stations & trains from data/trains final data .xlsx
        excel_path = os.path.join(DATA_DIR, "trains final data .xlsx")
        new_train_data_path = os.path.join(DATA_DIR, "NEW_TRAIN_DATA.xlsx")

        if os.path.exists(excel_path):
            wb = openpyxl.load_workbook(excel_path, data_only=True)
            for sheet_name in wb.sheetnames:
                if sheet_name.endswith("-2"):
                    continue # Skip duplicates
                
                sheet = wb[sheet_name]
                rows = list(sheet.iter_rows(values_only=True))
                if not rows or len(rows) < 2:
                    continue
                
                header = [str(c).strip() if c else "" for c in rows[0]]
                # Map column indices
                col_map = {name: idx for idx, name in enumerate(header)}
                
                train_num = str(sheet_name).replace("-2", "").strip()
                first_row = rows[1]
                
                t_name = str(first_row[col_map.get("Train_Name", 2)] or f"Express {train_num}")
                t_src = str(first_row[col_map.get("Source_Station", 3)] or "NDLS")
                t_dst = str(first_row[col_map.get("Destination_Station", 4)] or "ASR")

                # Upsert Train
                train_obj = db.query(Train).filter(Train.train_number == train_num).first()
                if not train_obj:
                    train_obj = Train(
                        train_number=train_num,
                        train_name=t_name,
                        train_type="Shatabdi" if "Shatabdi" in t_name else "Superfast Express",
                        source=f"{t_src}",
                        destination=f"{t_dst}",
                        total_distance_km=448.0 if "ASR" in t_dst else (303.0 if "KLK" in t_dst else 350.0),
                        is_active=True
                    )
                    db.add(train_obj)
                    db.flush()

                # Process stops for this train
                seen_stops = set()
                curr_dist = 0.0
                
                for r in rows[1:]:
                    if not r or len(r) <= col_map.get("Station_Code", 6):
                        continue
                    
                    st_code = str(r[col_map.get("Station_Code", 6)] or "").strip().replace("[", "").replace("]", "")
                    st_name = str(r[col_map.get("Station_Name", 7)] or st_code).strip()
                    if not st_code or st_code in seen_stops:
                        continue
                    seen_stops.add(st_code)

                    # Station coordinates lookup
                    info = STATION_COORDS.get(st_code, {
                        "lat": 28.6 + (len(seen_stops) * 0.2),
                        "lng": 77.2 - (len(seen_stops) * 0.15),
                        "name": st_name,
                        "zone": "NR",
                        "div": "DLI"
                    })

                    st_obj = db.query(Station).filter(Station.code == st_code).first()
                    if not st_obj:
                        st_obj = Station(
                            code=st_code,
                            name=info.get("name", st_name),
                            latitude=info["lat"],
                            longitude=info["lng"],
                            zone=info.get("zone", "NR"),
                            division=info.get("div", "DLI")
                        )
                        db.add(st_obj)
                        db.flush()

                    seq = int(float(r[col_map.get("Station_Sequence", 8)] or len(seen_stops)))
                    arr = clean_time(r[col_map.get("Scheduled_Arrival", 9)])
                    dep = clean_time(r[col_map.get("Scheduled_Departure", 10)])
                    halt = int(float(r[col_map.get("Scheduled_Halt_Minutes", 11)] or 2))
                    
                    # Estimate progressive distance
                    curr_dist = round((seq - 1) * 48.5, 1)

                    # Upsert RouteStop
                    stop_obj = db.query(RouteStop).filter(
                        RouteStop.train_id == train_obj.id,
                        RouteStop.station_id == st_obj.id
                    ).first()
                    
                    if not stop_obj:
                        db.add(RouteStop(
                            train_id=train_obj.id,
                            station_id=st_obj.id,
                            stop_sequence=seq,
                            scheduled_arrival=arr,
                            scheduled_departure=dep,
                            distance_from_source_km=curr_dist,
                            halt_minutes=halt,
                            platform=str(min(5, (seq % 4) + 1))
                        ))

                    # Ingest speed restrictions & signal halt field observations from the dataset!
                    sig_halt = r[col_map.get("Signal_Halt_Occurred", 14)] if "Signal_Halt_Occurred" in col_map else False
                    sig_dur = float(r[col_map.get("Signal_Halt_Duration_Min", 15)] or 0) if "Signal_Halt_Duration_Min" in col_map else 0
                    sp_res = r[col_map.get("Speed_Restriction_Active", 17)] if "Speed_Restriction_Active" in col_map else False
                    res_reason = r[col_map.get("Restriction_Reason", 21)] if "Restriction_Reason" in col_map else None

                    if (sig_halt or sig_dur > 0) and st_code != "NDLS":
                        obs_desc = f"Historical signal halt recorded at {st_name} ({sig_dur:.0f} min hold)."
                        if not db.query(FieldObservation).filter(FieldObservation.description == obs_desc).first():
                            db.add(FieldObservation(
                                station_id=st_obj.id,
                                track_section=f"LINE-{st_code}",
                                pole_reference=f"KM {seq * 42}/10",
                                observation_type="SIGNAL_INTERLOCKING",
                                description=obs_desc,
                                impact_delay_min=int(sig_dur) if sig_dur > 0 else 6,
                                severity="HIGH" if sig_dur > 15 else "MEDIUM"
                            ))

                    if sp_res and res_reason:
                        obs_desc = f"Speed restriction active near {st_name}: {res_reason}"
                        if not db.query(FieldObservation).filter(FieldObservation.description == obs_desc).first():
                            db.add(FieldObservation(
                                station_id=st_obj.id,
                                track_section=f"LINE-{st_code}",
                                pole_reference=f"KM {seq * 38}/04",
                                observation_type="SPEED_RESTRICTION",
                                description=obs_desc,
                                impact_delay_min=5,
                                severity="MEDIUM"
                            ))

                db.commit()
            print("✓ Ingested all train routes and stops from data/trains final data .xlsx.")

        # 3. Ingest latest 6-month live telemetry log for 12497 Shane Punjab
        csv_log_path = os.path.join(DATA_DIR, "ruhani_12497_shan_e_punjab_6months.csv")
        if os.path.exists(csv_log_path):
            with open(csv_log_path, mode='r', encoding='utf-8', errors='ignore') as f:
                reader = csv.DictReader(f)
                rows = list(reader)
                if rows:
                    latest = rows[-1] # Most recent live telemetry row
                    t12497 = db.query(Train).filter(Train.train_number == "12497").first()
                    if t12497:
                        # Find next station
                        snp_st = db.query(Station).filter(Station.code == "SNP").first()
                        pnp_st = db.query(Station).filter(Station.code == "PNP").first()

                        state_obj = db.query(LiveTrainState).filter(LiveTrainState.train_id == t12497.id).first()
                        if not state_obj:
                            state_obj = LiveTrainState(train_id=t12497.id)
                            db.add(state_obj)

                        state_obj.current_station_id = snp_st.id if snp_st else None
                        state_obj.next_station_id = pnp_st.id if pnp_st else None
                        state_obj.latitude = float(latest.get("Latitude", 29.1850))
                        state_obj.longitude = float(latest.get("Longitude", 76.9920))
                        state_obj.speed_kmh = float(latest.get("Current_Speed", 85.0))
                        state_obj.current_delay_min = int(float(latest.get("Current_Delay", 14)))
                        state_obj.remaining_distance_km = 379.0
                        state_obj.track_section = "NDLS-PNP-LINE1"
                        state_obj.reference_pole = "KM 68/14"
                        state_obj.status = "RUNNING"
                        state_obj.last_updated = datetime.utcnow()
                        db.commit()
                        print(f"✓ Initialized Train 12497 telemetry from {csv_log_path}: Speed {state_obj.speed_kmh} km/h, Delay +{state_obj.current_delay_min}m.")

        # 4. Ingest live state for Kalka Shatabdi 12011 as following rake
        t12011 = db.query(Train).filter(Train.train_number == "12011").first()
        if t12011:
            ndls_st = db.query(Station).filter(Station.code == "NDLS").first()
            pnp_st = db.query(Station).filter(Station.code == "PNP").first()
            state_12011 = db.query(LiveTrainState).filter(LiveTrainState.train_id == t12011.id).first()
            if not state_12011:
                state_12011 = LiveTrainState(
                    train_id=t12011.id,
                    current_station_id=ndls_st.id if ndls_st else None,
                    next_station_id=pnp_st.id if pnp_st else None,
                    latitude=28.8200,
                    longitude=77.1200,
                    speed_kmh=92.0,
                    current_delay_min=6,
                    remaining_distance_km=271.0,
                    track_section="NDLS-PNP-LINE1",
                    reference_pole="KM 28/06",
                    status="RUNNING",
                    last_updated=datetime.utcnow()
                )
                db.add(state_12011)
                db.commit()

        print("=== DATA INGESTION COMPLETE ===")
    finally:
        db.close()

if __name__ == "__main__":
    load_all_data()
